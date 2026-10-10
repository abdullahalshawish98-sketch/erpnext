"""
مراجعة كروت الوكيل "عدنان" فقط (لا علاقة لها بمتجر سلتك) — نفس منطق
detailed_card_review.py (مطابقة مباشرة -> إقران استرجاع 10-180 دقيقة ->
إقران موسَّع حتى 24 ساعة)، لكن الحركات هنا مُدخَلة يدوياً من سكرين شوت
تطبيق شي ان (لا يوجد ملف تصدير مُجمَّع لهذه الكروت)، وكل سكرين شوت راجعه
المستخدم بنفسه. قراءة فقط — لا يُعدِّل أي بيانات في ERPNext.

الاستخدام: python scripts/adnan_cards_review.py [مسار ملف الإخراج xlsx]
"""
import sys
import json
import re
from datetime import datetime
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill

out_path = sys.argv[1] if len(sys.argv) > 1 else 'data/adnan_cards_review.xlsx'
codes_file = '/root/.claude/uploads/07a21562-e623-57de-a7c6-ab6649a6aa2d/0b3984df-Untitled_spreadsheet.xlsx'

# ---------------------------------------------------------------------------
# حركات شي ان المُدخَلة يدوياً من سكرين شوت راجعها المستخدم حركة حركة.
# نفس اصطلاح الإشارة المستخدَم في detailed_card_review.py: الشراء بقيمة
# سالبة، الاسترجاع والتفعيل بقيمة موجبة.
# ---------------------------------------------------------------------------
MANUAL_CARD_EVENTS = {
    "1118717901813003296": [
        ("2026-09-23", "18:35", "H********2@gmail.com", "Activated", 300),
        ("2026-09-28", "02:26", "j****************2@hotm...", "Purchase", -42.33),
        ("2026-09-28", "02:41", "m****************0@hot...", "Purchase", -55.32),
        ("2026-09-28", "02:42", "j****************2@hotm...", "Purchase", -41.98),
        ("2026-09-28", "02:56", "m****************0@hot...", "Purchase", -85.62),
        ("2026-09-28", "03:06", "m****************0@hot...", "Purchase", -52.27),
        ("2026-09-28", "03:42", "m****************0@hot...", "Refund", 26.17),
        ("2026-09-28", "03:42", "j****************2@hotm...", "Refund", 27.77),
        ("2026-09-28", "03:57", "m****************0@hot...", "Refund", 8.35),
        ("2026-09-28", "04:06", "m****************0@hot...", "Refund", 28.49),
        ("2026-09-29", "22:00", "a****************2@hot...", "Purchase", -44.38),
        ("2026-09-29", "22:07", "a****************2@hot...", "Purchase", -42.80),
        ("2026-09-29", "23:00", "a****************2@hot...", "Refund", 23.51),
        ("2026-09-29", "23:07", "a****************2@hot...", "Refund", 31.27),
        ("2026-09-29", "23:14", "m****************1@ho...", "Purchase", -44.52),
        ("2026-09-30", "00:15", "m****************1@ho...", "Refund", 18.43),
        ("2026-09-30", "00:44", "h****************8@hotmail...", "Purchase", -41.29),
        ("2026-09-30", "01:44", "h****************8@hotmail...", "Refund", 21.31),
    ],
    "1119717811679800390": [
        ("2026-06-11", "10:53", "h***********7@gmail.com", "Activated", 500),
        ("2026-08-10", "06:39", "r***********t@hotmail.com", "Purchase", -47.56),
        ("2026-08-10", "06:47", "r***********t@hotmail.com", "Purchase", -49.66),
        ("2026-08-10", "06:51", "r***********t@hotmail.com", "Purchase", -42.60),
        ("2026-08-10", "06:56", "r***********t@hotmail.com", "Purchase", -52.33),
        ("2026-08-10", "06:57", "m****************k@hot...", "Purchase", -55.28),
        ("2026-08-10", "07:07", "j***********g@hotmail.com", "Purchase", -45.11),
        ("2026-08-10", "07:07", "m****************k@hot...", "Purchase", -45.04),
        ("2026-08-10", "07:10", "j***********g@hotmail.com", "Purchase", -50.87),
        ("2026-08-10", "07:20", "j***********g@hotmail.com", "Purchase", -43.15),
        ("2026-08-10", "07:24", "j***********g@hotmail.com", "Purchase", -44.43),
        ("2026-08-10", "07:40", "r***********t@hotmail.com", "Refund", 30.43),
        ("2026-08-10", "07:47", "r***********t@hotmail.com", "Refund", 21.13),
        ("2026-08-10", "07:51", "r***********t@hotmail.com", "Refund", 30.63),
        ("2026-08-10", "08:07", "j***********g@hotmail.com", "Refund", 30.41),
        ("2026-08-10", "08:07", "m****************k@hot...", "Refund", 9.80),
        ("2026-08-10", "08:11", "j***********g@hotmail.com", "Refund", 12.50),
        ("2026-08-10", "08:21", "j***********g@hotmail.com", "Refund", 15.78),
        ("2026-08-10", "08:24", "j***********g@hotmail.com", "Refund", 30.87),
        ("2026-08-11", "02:13", "c**************l@hotmail.c...", "Purchase", -43.04),
        ("2026-08-11", "02:45", "c**************l@hotmail.c...", "Purchase", -45.79),
        ("2026-08-11", "02:54", "c**************l@hotmail.c...", "Purchase", -41.40),
        ("2026-08-11", "03:01", "c**************l@hotmail.c...", "Purchase", -43.79),
        ("2026-08-11", "03:14", "c**************l@hotmail.c...", "Refund", 8.95),
        ("2026-08-11", "03:54", "c**************l@hotmail.c...", "Refund", 29.94),
        ("2026-08-11", "04:01", "c**************l@hotmail.c...", "Refund", 13.88),
        ("2026-08-12", "13:40", "d***************k@hotmail...", "Purchase", -49.29),
    ],
    "1113717905015800023": [
        ("2026-09-27", "11:33", "H********2@gmail.com", "Activated", 300),
        ("2026-10-02", "06:26", "l***************1@hotmail...", "Purchase", -42.91),
        ("2026-10-02", "06:38", "l***************1@hotmail...", "Purchase", -46.04),
        ("2026-10-02", "06:42", "l***************1@hotmail...", "Purchase", -47.49),
        ("2026-10-02", "07:26", "l***************1@hotmail...", "Refund", 14.03),
        ("2026-10-02", "07:39", "l***************1@hotmail...", "Refund", 30.52),
        ("2026-10-02", "07:43", "l***************1@hotmail...", "Refund", 12.14),
        ("2026-10-02", "18:32", "p*******************5@hot...", "Purchase", -41.81),
        ("2026-10-02", "18:38", "p*******************5@hot...", "Purchase", -42.66),
        ("2026-10-02", "19:00", "p*******************5@hot...", "Purchase", -45.16),
        ("2026-10-02", "19:07", "j*****************4@hotm...", "Purchase", -56.68),
        ("2026-10-02", "19:39", "p*******************5@hot...", "Refund", 14.70),
        ("2026-10-02", "20:00", "p*******************5@hot...", "Refund", 25.61),
        ("2026-10-05", "03:36", "e*************4@outlook.c...", "Purchase", -46.99),
        ("2026-10-05", "04:36", "e*************4@outlook.c...", "Refund", 31.40),
        ("2026-10-08", "03:53", "a***************2@outlook...", "Purchase", -42.30),
        ("2026-10-08", "04:54", "a***************2@outlook...", "Refund", 25.21),
    ],
    "1114017882943009139": [
        ("2026-09-01", "22:25", "P***e@yahoo.com", "Activated", 500),
        ("2026-09-03", "04:52", "e************6@hotmail.c...", "Purchase", -42.68),
        ("2026-09-03", "05:16", "j*****************4@hotm...", "Purchase", -42.25),
        ("2026-09-03", "05:52", "e************6@hotmail.c...", "Refund", 16.52),
        ("2026-09-03", "06:17", "j*****************4@hotm...", "Refund", 10.89),
        ("2026-09-03", "15:38", "j*****************4@hotm...", "Purchase", -40.71),
        ("2026-09-03", "15:50", "j*****************4@hotm...", "Purchase", -42.41),
        ("2026-09-03", "15:55", "j*****************4@hotm...", "Purchase", -40.96),
        ("2026-09-03", "16:03", "e************6@hotmail.c...", "Purchase", -49.14),
        ("2026-09-03", "16:38", "j*****************4@hotm...", "Refund", 26.94),
        ("2026-09-03", "16:51", "j*****************4@hotm...", "Refund", 10.89),
        ("2026-09-03", "16:56", "j*****************4@hotm...", "Refund", 12.58),
        ("2026-09-03", "17:03", "e************6@hotmail.c...", "Refund", 8.71),
        ("2026-09-04", "12:13", "k**************0@hotmail...", "Purchase", -42.56),
        ("2026-09-04", "12:49", "k**************0@hotmail...", "Purchase", -43.66),
        ("2026-09-04", "13:11", "k**************0@hotmail...", "Purchase", -67.57),
        ("2026-09-04", "13:14", "k**************0@hotmail...", "Refund", 21.13),
        ("2026-09-04", "13:50", "k**************0@hotmail...", "Refund", 31.43),
        ("2026-09-06", "03:49", "c******************0@hot...", "Purchase", -50.13),
        ("2026-09-06", "04:00", "c******************0@hot...", "Purchase", -40.90),
        ("2026-09-06", "04:04", "c******************0@hot...", "Purchase", -41.77),
        ("2026-09-06", "04:18", "c******************0@hot...", "Purchase", -41.79),
        ("2026-09-06", "04:24", "c******************0@hot...", "Purchase", -41.19),
        ("2026-09-06", "05:00", "c******************0@hot...", "Refund", 12.17),
        ("2026-09-06", "05:18", "c******************0@hot...", "Refund", 22.74),
        ("2026-09-06", "05:24", "c******************0@hot...", "Refund", 10.10),
        ("2026-09-07", "10:00", "j**************3@hotmail.c...", "Purchase", -42.55),
        ("2026-09-07", "11:01", "j**************3@hotmail.c...", "Refund", 25.32),
    ],
    "1116017867995202786": [
        ("2026-08-15", "15:12", "h*******n@gmail.com", "Activated", 500),
        ("2026-08-16", "05:09", "c***************g@hotmail...", "Purchase", -46.90),
        ("2026-08-16", "05:23", "m*************d@hotmail...", "Purchase", -61.41),
        ("2026-08-16", "05:45", "m*************d@hotmail...", "Purchase", -97.78),
        ("2026-08-16", "06:10", "c***************g@hotmail...", "Refund", 9.18),
        ("2026-08-16", "06:24", "m*************d@hotmail...", "Refund", 13.66),
        ("2026-08-17", "00:10", "m*************d@hotmail...", "Purchase", -48.19),
        ("2026-08-17", "00:18", "m*************d@hotmail...", "Purchase", -42.21),
        ("2026-08-17", "00:31", "r***********w@hotmail.com", "Purchase", -41.89),
        ("2026-08-17", "00:44", "r***********w@hotmail.com", "Purchase", -43.85),
        ("2026-08-17", "01:11", "m*************d@hotmail...", "Refund", 22.74),
        ("2026-08-17", "01:32", "r***********w@hotmail.com", "Refund", 27.33),
        ("2026-08-18", "02:54", "j*******************7@hot...", "Purchase", -41.85),
        ("2026-08-18", "03:04", "j*******************7@hot...", "Purchase", -42.99),
        ("2026-08-18", "03:12", "j*******************7@hot...", "Purchase", -43.50),
        ("2026-08-18", "03:47", "j*******************7@hot...", "Purchase", -40.81),
        ("2026-08-18", "03:54", "j*******************7@hot...", "Refund", 22.71),
        ("2026-08-18", "04:04", "j*******************7@hot...", "Refund", 28.83),
        ("2026-08-18", "04:12", "j*******************7@hot...", "Refund", 14.70),
        ("2026-08-18", "04:27", "f****************1@hotm...", "Purchase", -48.24),
        ("2026-08-18", "04:48", "j*******************7@hot...", "Refund", 13.61),
        ("2026-08-18", "05:28", "f****************1@hotm...", "Refund", 9.75),
    ],
    "1116717907666003070": [
        ("2026-09-30", "13:10", "H********2@gmail.com", "Activated", 500),
        ("2026-10-02", "19:37", "m****************5@hotm...", "Purchase", -43.74),
        ("2026-10-02", "19:41", "m****************5@hotm...", "Purchase", -42.62),
        ("2026-10-02", "20:37", "m****************5@hotm...", "Refund", 22.90),
        ("2026-10-02", "20:41", "m****************5@hotm...", "Refund", 30.98),
        ("2026-10-03", "00:50", "m****************5@hotm...", "Purchase", -42.64),
        ("2026-10-03", "03:44", "w****************2@hot...", "Purchase", -44.86),
        ("2026-10-03", "03:50", "w****************2@hot...", "Purchase", -42.31),
        ("2026-10-03", "03:59", "w****************2@hot...", "Purchase", -40.74),
        ("2026-10-03", "04:09", "w****************2@hot...", "Purchase", -45.03),
        ("2026-10-03", "04:21", "j***************6@hotmail...", "Purchase", -41.20),
        ("2026-10-03", "04:28", "j***************6@hotmail...", "Purchase", -40.71),
        ("2026-10-03", "04:34", "j***************6@hotmail...", "Purchase", -42.21),
        ("2026-10-03", "04:39", "j***************6@hotmail...", "Purchase", -44.25),
        ("2026-10-03", "04:44", "w****************2@hot...", "Refund", 30.51),
        ("2026-10-03", "04:50", "w****************2@hot...", "Refund", 27.54),
        ("2026-10-03", "04:51", "c****************3@hotm...", "Purchase", -40.60),
        ("2026-10-03", "04:56", "c****************3@hotm...", "Purchase", -41.36),
        ("2026-10-03", "04:58", "c****************3@hotm...", "Purchase", -47.53),
        ("2026-10-03", "05:00", "w****************2@hot...", "Refund", 20.83),
        ("2026-10-03", "05:09", "w****************2@hot...", "Refund", 8.76),
        ("2026-10-03", "05:17", "b***************7@hotmail....", "Purchase", -40.57),
        ("2026-10-03", "05:22", "j***************6@hotmail...", "Refund", 25.21),
        ("2026-10-03", "05:28", "j***************6@hotmail...", "Refund", 30.17),
        ("2026-10-03", "05:33", "b***************7@hotmail....", "Purchase", -47.81),
        ("2026-10-03", "05:34", "j***************6@hotmail...", "Refund", 8.56),
        ("2026-10-03", "05:39", "j***************6@hotmail...", "Refund", 15.82),
        ("2026-10-03", "05:51", "c****************3@hotm...", "Refund", 19.25),
        ("2026-10-03", "05:56", "c****************3@hotm...", "Refund", 25.61),
        ("2026-10-03", "05:59", "c****************3@hotm...", "Refund", 10.54),
        ("2026-10-03", "06:17", "b***************7@hotmail....", "Refund", 13.68),
        ("2026-10-03", "06:33", "b***************7@hotmail....", "Refund", 5.76),
        ("2026-10-03", "06:40", "b***************7@hotmail....", "Purchase", -58.67),
        ("2026-10-03", "07:40", "b***************7@hotmail....", "Refund", 10.16),
    ],
}

# كروت "مربوطة": آلية شي ان تسمح بالشراء من بطاقة واحدة بـ5 حسابات فقط؛ عند
# الشراء بحساب سادس يُطلَب ربط البطاقة بذلك الحساب، والمستخدم لا يملك دخولاً
# له فلا يمكنه فتح تقرير شي ان الكامل لهذه الكروت. الحساب السادس مُستخرَج من
# حقل Shopping Account في آخر فاتورة شراء لكل بطاقة.
LINKED_CARDS = [
    {
        'card': '1116017910192605944',
        'shopping_account': 'LarryHoldenBXm36511@outlook.com - شي ان',
        'note': 'مربوط - لا يمكن الوصول لتقرير شي ان إلا بفتح هذا الحساب',
    },
    {
        'card': '1118017907661203049',
        'shopping_account': 'LauraBlackF22871@hotmail.com - شي ان',
        'note': 'مربوط - لا يمكن الوصول لتقرير شي ان إلا بفتح هذا الحساب (آخر فحص آلي من سكرين شوت جزئي أظهر فرق 4.88$، 4.00$ منها فاتورة ACC-PINV-2026-24843 بقيمة مختلفة عن الشراء المقابل تماماً بالتوقيت — بانتظار التأكيد من الحساب السادس)',
    },
    {
        'card': '1116217899822802855',
        'shopping_account': 'LarryHoldenBXm36511@outlook.com - شي ان',
        'note': 'مربوط (نفس حساب الكرت 1116017910192605944) - تم تجاوزه سابقاً',
    },
    {
        'card': '1113117819597404308',
        'shopping_account': 'ledgercostavhml@outlook.com - شي ان',
        'note': 'مربوط - لا يمكن الوصول لتقرير شي ان إلا بفتح هذا الحساب',
    },
    {
        'card': '1117117809261204915',
        'shopping_account': 'adlerlindseyjaos@outlook.com - شي ان',
        'note': 'مربوط - لا يمكن الوصول لتقرير شي ان إلا بفتح هذا الحساب',
    },
    {
        'card': '1118417837961608167',
        'shopping_account': 'HyacinthaCarreon340@hotmail.com - شي ان',
        'note': 'مربوط - لا يمكن الوصول لتقرير شي ان إلا بفتح هذا الحساب',
    },
]

# كروت راجعها المستخدم ووجد فروقها مقبولة (تقريب بسيط فقط، لا خطأ مالي
# حقيقي) — تُستبعَد من ورقة "غير المطابقة" مثل USER_CONFIRMED_CLEAN في
# detailed_card_review.py.
USER_CONFIRMED_CLEAN = {
    "1113717905015800023",  # أكّد المستخدم 2026-10-10: مطابق، تجاهل فروقات الفواتير البسيطة (0.40$)
    "1116717907666003070",  # مطابق آلياً — فروق تقريب ضئيلة فقط (أكبرها 0.47$)، لا خطأ مالي
}

with open('data/Gift_Card.json', encoding='utf-8') as f:
    cards = json.load(f)
with open('data/Purchase_Invoice.json', encoding='utf-8') as f:
    pis = json.load(f)


def base_serial(name):
    m = re.search(r'(\d{10,})', name)
    return m.group(1) if m else None


card_by_serial = {}
card_record_by_serial = {}
for c in cards:
    bs = base_serial(c['name'])
    if bs:
        card_by_serial[bs] = c['name']
        card_record_by_serial[bs] = c

code_by_serial = {}
try:
    wb_codes = openpyxl.load_workbook(codes_file, data_only=True)
    ws_codes = wb_codes['Sheet1']
    for row in ws_codes.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        serial, code = str(row[0]).strip(), row[1]
        if serial and code is not None:
            code_by_serial[serial] = int(code) if isinstance(code, float) and code.is_integer() else code
except FileNotFoundError:
    pass


def dt(date_, time_):
    return datetime.fromisoformat(f"{date_} {time_}")


def inv_dt(p):
    t = str(p.get('posting_time') or '00:00:00').split('.')[0]
    hh, mm, ss = (t.split(':') + ['0', '0', '0'])[:3]
    return datetime.fromisoformat(f"{p.get('posting_date')} {int(hh):02d}:{int(mm):02d}:{int(float(ss)):02d}")


def take(erp_by_value, value, near_dt=None, max_gap_hours=None):
    v = round(value, 2)
    if v in erp_by_value:
        key = v
    else:
        close_keys = [k for k in erp_by_value if abs(k - v) <= 0.02 + 1e-9]
        key = min(close_keys, key=lambda k: abs(k - v)) if close_keys else None
    if key is None:
        return None
    entries = erp_by_value.get(key)
    if not entries:
        return None
    if near_dt is None:
        return entries.pop()[0]
    best_idx, best_gap = None, None
    for i, (_, idt) in enumerate(entries):
        gap = abs((idt - near_dt).total_seconds())
        if best_gap is None or gap < best_gap:
            best_idx, best_gap = i, gap
    if max_gap_hours is not None and best_gap > max_gap_hours * 3600:
        return None
    return entries.pop(best_idx)[0]


def candidates(purchases, refunds, p_idx, r_idx, lo_min, hi_min, target_min):
    cands = []
    for pi in p_idx:
        for ri in r_idx:
            delta_min = (dt(refunds[ri]['date'], refunds[ri]['time']) - dt(purchases[pi]['date'], purchases[pi]['time'])).total_seconds() / 60
            if lo_min <= delta_min <= hi_min:
                cands.append((abs(delta_min - target_min), pi, ri))
    cands.sort(key=lambda x: x[0])
    return cands


def process_card(cid, events_all):
    name = card_by_serial.get(cid)
    erp_rows = sorted([p for p in pis if p.get('gift_card') == name and p.get('docstatus') == 1],
                       key=lambda p: (p.get('posting_date', ''), p.get('creation', '')))
    erp_by_value = defaultdict(list)
    for p in erp_rows:
        erp_by_value[round(p.get('total') or 0, 2)].append((p.get('name'), inv_dt(p)))
    all_invoice_names = {p.get('name') for p in erp_rows}

    detail_rows = []
    used_invoice_names = set()

    activated = [t for t in events_all if t[3] == 'Activated']
    for t in activated:
        detail_rows.append({
            'card': cid, 'date': t[0], 'time': t[1], 'account': t[2],
            'type': 'تفعيل', 'shein_amount': t[4], 'invoice': '', 'invoice_value': '',
            'status': 'تفعيل', 'note': '',
        })

    accounts = sorted({t[2] for t in events_all if t[3] != 'Activated'})
    acc_all_purchase_results = defaultdict(list)

    for acc in accounts:
        events = sorted((t for t in events_all if t[3] != 'Activated' and t[2] == acc), key=lambda t: (t[0], t[1]))
        sessions = []
        cur, prev_dt = [], None
        for t in events:
            cur_dt = dt(t[0], t[1])
            if prev_dt and (cur_dt - prev_dt).total_seconds() > 3 * 3600:
                sessions.append(cur)
                cur = []
            cur.append(t)
            prev_dt = cur_dt
        if cur:
            sessions.append(cur)

        acc_leftover_purchases = []
        acc_leftover_refunds = []

        for sess in sessions:
            purchases = [{'amount': -t[4], 'date': t[0], 'time': t[1], 'account': acc} for t in sess if t[3] == 'Purchase']
            refunds = [{'amount': t[4], 'date': t[0], 'time': t[1], 'account': acc} for t in sess if t[3] == 'Refund']

            if purchases and refunds and abs(sum(p['amount'] for p in purchases) - sum(r['amount'] for r in refunds)) < 0.05:
                for p in purchases:
                    detail_rows.append({
                        'card': cid, 'date': p['date'], 'time': p['time'], 'account': acc,
                        'type': 'شراء', 'shein_amount': -p['amount'], 'invoice': '', 'invoice_value': '',
                        'status': '✅ جلسة مُلغاة بالكامل', 'note': 'مجموع الشراء = مجموع الاسترجاع',
                    })
                for r in refunds:
                    detail_rows.append({
                        'card': cid, 'date': r['date'], 'time': r['time'], 'account': acc,
                        'type': 'استرجاع', 'shein_amount': r['amount'], 'invoice': '', 'invoice_value': '',
                        'status': '✅ جلسة مُلغاة بالكامل', 'note': 'ضمن إلغاء الجلسة الكاملة',
                    })
                continue

            all_p, all_r = list(range(len(purchases))), list(range(len(refunds)))
            used_p, used_r = set(), set()
            note_of = {}

            for _, pi, ri in candidates(purchases, refunds, all_p, all_r, 10, 180, 60):
                if pi in used_p or ri in used_r:
                    continue
                net = round(purchases[pi]['amount'] - refunds[ri]['amount'], 2)
                if abs(net) < 0.05:
                    used_p.add(pi)
                    used_r.add(ri)
                    gap = (dt(refunds[ri]['date'], refunds[ri]['time']) - dt(purchases[pi]['date'], purchases[pi]['time'])).total_seconds() / 60
                    note_of[pi] = (f"أُلغيت بالكامل باسترجاع +{refunds[ri]['amount']}$ بعد {gap:.0f} دقيقة (صافي صفر)", 0.0, None)
                    detail_rows.append({
                        'card': cid, 'date': refunds[ri]['date'], 'time': refunds[ri]['time'], 'account': acc,
                        'type': 'استرجاع', 'shein_amount': refunds[ri]['amount'], 'invoice': '', 'invoice_value': '',
                        'status': '✅ إلغاء كامل (صافي صفر)', 'note': f"يقابل شراء {purchases[pi]['date']} {purchases[pi]['time']}",
                    })
                    continue
                matched_name = take(erp_by_value, net, near_dt=dt(purchases[pi]['date'], purchases[pi]['time']))
                if not matched_name:
                    continue
                used_p.add(pi)
                used_r.add(ri)
                gap = (dt(refunds[ri]['date'], refunds[ri]['time']) - dt(purchases[pi]['date'], purchases[pi]['time'])).total_seconds() / 60
                note_of[pi] = (f"مقرون باسترجاع +{refunds[ri]['amount']}$ بعد {gap:.0f} دقيقة", net, matched_name)
                detail_rows.append({
                    'card': cid, 'date': refunds[ri]['date'], 'time': refunds[ri]['time'], 'account': acc,
                    'type': 'استرجاع', 'shein_amount': refunds[ri]['amount'],
                    'invoice': matched_name, 'invoice_value': '',
                    'status': '✅ مطروح من فاتورة الشراء المقابلة',
                    'note': f"نُسخ ضمن صافي الشراء {purchases[pi]['date']} {purchases[pi]['time']}",
                })

            remaining_p = [i for i in all_p if i not in used_p]
            for pi in remaining_p:
                matched_name = take(erp_by_value, purchases[pi]['amount'],
                                     near_dt=dt(purchases[pi]['date'], purchases[pi]['time']))
                if matched_name:
                    used_p.add(pi)
                    note_of[pi] = ('مطابقة مباشرة بقيمة شي ان الكاملة', purchases[pi]['amount'], matched_name)

            for pi, p in enumerate(purchases):
                if pi in note_of:
                    note, net, matched_name = note_of[pi]
                else:
                    note, net, matched_name = 'بلا استرجاع (بعد)', round(p['amount'], 2), None
                p_result = {**p, 'net': net, 'note': note, 'match': matched_name}
                acc_all_purchase_results[acc].append(p_result)
                if pi not in note_of:
                    acc_leftover_purchases.append(p_result)

            for ri in all_r:
                if ri not in used_r:
                    acc_leftover_refunds.append(refunds[ri])

        all_p2, all_r2 = list(range(len(acc_leftover_purchases))), list(range(len(acc_leftover_refunds)))
        used_p2, used_r2 = set(), set()
        for _, pi, ri in candidates(acc_leftover_purchases, acc_leftover_refunds, all_p2, all_r2, 180, 1440, 180):
            if pi in used_p2 or ri in used_r2:
                continue
            net = round(acc_leftover_purchases[pi]['net'] - acc_leftover_refunds[ri]['amount'], 2)
            matched_name = take(erp_by_value, net, near_dt=dt(acc_leftover_purchases[pi]['date'], acc_leftover_purchases[pi]['time']))
            if matched_name:
                used_p2.add(pi)
                used_r2.add(ri)
                r = acc_leftover_refunds[ri]
                gap = (dt(r['date'], r['time']) - dt(acc_leftover_purchases[pi]['date'], acc_leftover_purchases[pi]['time'])).total_seconds() / 60
                acc_leftover_purchases[pi]['net'] = net
                acc_leftover_purchases[pi]['note'] = f"⚠ شاذة — مقرون باسترجاع متأخر {r['date']} {r['time']} بعد {gap:.0f} دقيقة"
                acc_leftover_purchases[pi]['match'] = matched_name
                detail_rows.append({
                    'card': cid, 'date': r['date'], 'time': r['time'], 'account': acc,
                    'type': 'استرجاع', 'shein_amount': r['amount'], 'invoice': matched_name or '', 'invoice_value': '',
                    'status': '✅ مطروح (إقران موسَّع حتى 24 ساعة)',
                    'note': f"نُسخ ضمن صافي الشراء {acc_leftover_purchases[pi]['date']} {acc_leftover_purchases[pi]['time']}",
                })

        for r_idx2, r in enumerate(acc_leftover_refunds):
            if r_idx2 in used_r2:
                continue
            detail_rows.append({
                'card': cid, 'date': r['date'], 'time': r['time'], 'account': r['account'],
                'type': 'استرجاع', 'shein_amount': r['amount'], 'invoice': '', 'invoice_value': '',
                'status': '⚠ استرجاع غير مُستخدَم', 'note': '',
            })

        for p in acc_all_purchase_results[acc]:
            if 'match' not in p:
                continue
            match = p.get('match')
            net = p['net']
            if match:
                used_invoice_names.add(match)
                inv_val = next((pp.get('total') for pp in erp_rows if pp.get('name') == match), '')
                status = f'✅ مطابق {match}'
            elif abs(net) < 0.05:
                status = '✅ إلغاء كامل (صافي صفر)'
                inv_val = ''
            else:
                status = '❌ غير مفسَّر — لا توجد فاتورة بهذه القيمة الصافية'
                inv_val = ''
            detail_rows.append({
                'card': cid, 'date': p['date'], 'time': p['time'], 'account': p['account'],
                'type': 'شراء', 'shein_amount': -p['amount'] if 'amount' in p else '',
                'invoice': match or '', 'invoice_value': inv_val,
                'status': status, 'note': f"{p['note']} | صافي متوقَّع={net}$",
            })

    unused_invoices = all_invoice_names - used_invoice_names
    for inv_name in unused_invoices:
        inv = next(p for p in erp_rows if p.get('name') == inv_name)
        detail_rows.append({
            'card': cid, 'date': inv.get('posting_date'), 'time': str(inv.get('posting_time'))[:8],
            'account': '', 'type': 'فاتورة مرتجع' if inv.get('is_return') else 'فاتورة شراء',
            'shein_amount': '', 'invoice': inv_name, 'invoice_value': inv.get('total'),
            'status': '🟠 فاتورة يتيمة — لا تقابلها أي حركة شي ان مطابقة', 'note': '',
        })

    detail_rows.sort(key=lambda r: (str(r['date']), str(r['time'])))
    code = code_by_serial.get(cid)
    for r in detail_rows:
        r['code'] = code

    real_balance = card_record_by_serial.get(cid, {}).get('balance')
    purchases_sum = sum(-t[4] for t in events_all if t[3] == 'Purchase')
    refunds_sum = sum(t[4] for t in events_all if t[3] == 'Refund')
    shein_real_balance = round((activated[0][4] if activated else 0) - purchases_sum + refunds_sum, 2)
    unmatched_purchases = sum(1 for r in detail_rows if r['type'] == 'شراء' and '❌' in r['status'])
    unused_refunds = sum(1 for r in detail_rows if r['status'] == '⚠ استرجاع غير مُستخدَم')
    orphan_invoices = sum(1 for r in detail_rows if '🟠' in r['status'])

    summary = {
        'card': cid, 'code': code, 'name': name, 'agent': 'عدنان',
        'activated': activated[0][4] if activated else '',
        'erpnext_balance': real_balance,
        'shein_real_balance': shein_real_balance,
        'diff': round((real_balance or 0) - shein_real_balance, 2) if real_balance is not None else '',
        'unmatched_purchases': unmatched_purchases,
        'unused_refunds': unused_refunds,
        'orphan_invoices': orphan_invoices,
    }
    return detail_rows, summary


all_detail_rows = []
all_summaries = []
for cid, events in MANUAL_CARD_EVENTS.items():
    if cid not in card_by_serial:
        print(f"!! تحذير: لا توجد بطاقة في ERPNext للسيريال {cid}")
        continue
    detail_rows, summary = process_card(cid, events)
    all_detail_rows.extend(detail_rows)
    all_summaries.append(summary)
    flag = '  <<< يحتاج مراجعة' if (summary['unmatched_purchases'] or summary['unused_refunds'] or summary['orphan_invoices'] or abs(summary['diff'] or 0) > 0.05) else ''
    print(f"{cid}  erp={summary['erpnext_balance']}  shein={summary['shein_real_balance']}  diff={summary['diff']}  "
          f"غير مطابق={summary['unmatched_purchases']}  استرجاع غير مستخدم={summary['unused_refunds']}  فواتير يتيمة={summary['orphan_invoices']}{flag}")

all_summaries.sort(key=lambda s: abs(s['diff'] if isinstance(s['diff'], (int, float)) else 0), reverse=True)

# ---------------------------------------------------------------------------
# بناء ملف Excel
# ---------------------------------------------------------------------------
wb = openpyxl.Workbook()
red = PatternFill('solid', fgColor='FFC7CE')
orange = PatternFill('solid', fgColor='FFE0B2')
green = PatternFill('solid', fgColor='C6EFCE')
grey = PatternFill('solid', fgColor='E0E0E0')
header_fill = PatternFill('solid', fgColor='4472C4')
header_font = Font(bold=True, color='FFFFFF')

LINKED_CARD_SERIALS = {c['card'] for c in LINKED_CARDS}


def is_clean_summary(s):
    if s['card'] in USER_CONFIRMED_CLEAN or s['card'] in LINKED_CARD_SERIALS:
        return True
    return s['unmatched_purchases'] == 0 and s['unused_refunds'] == 0 and s['orphan_invoices'] == 0 and abs(s['diff'] or 0) <= 0.05


unmatched_summaries = [s for s in all_summaries if not is_clean_summary(s)]
unmatched_card_ids = {s['card'] for s in unmatched_summaries}

ws1 = wb.active
ws1.title = 'الكروت غير المطابقة'
ws1.sheet_view.rightToLeft = True
headers1 = ['السيريال', 'الرمز', 'اسم البطاقة', 'الوكيل', 'التفعيل $', 'رصيد المنظومة $', 'الرصيد الحقيقي (شي ان) $',
            'الفرق $', 'مشتريات غير مفسَّرة', 'استرجاعات غير مستخدَمة', 'فواتير يتيمة']
for ci, h in enumerate(headers1, 1):
    c = ws1.cell(1, ci, h)
    c.font = header_font
    c.fill = header_fill
for ri, s in enumerate(unmatched_summaries, 2):
    ws1.cell(ri, 1, s['card'])
    ws1.cell(ri, 2, s['code'])
    ws1.cell(ri, 3, s['name'])
    ws1.cell(ri, 4, s['agent'])
    ws1.cell(ri, 5, s['activated'])
    ws1.cell(ri, 6, s['erpnext_balance'])
    ws1.cell(ri, 7, s['shein_real_balance'])
    ws1.cell(ri, 8, s['diff'])
    ws1.cell(ri, 9, s['unmatched_purchases'])
    ws1.cell(ri, 10, s['unused_refunds'])
    ws1.cell(ri, 11, s['orphan_invoices'])
    fill = red if abs(s['diff'] or 0) > 0.05 else orange
    for ci in range(1, 12):
        ws1.cell(ri, ci).fill = fill
for ci, w in zip(range(1, 12), [22, 10, 34, 12, 12, 16, 20, 10, 14, 16, 12]):
    ws1.column_dimensions[chr(64 + ci)].width = w

ws2 = wb.create_sheet('تفاصيل الكروت غير المطابقة')
ws2.sheet_view.rightToLeft = True
headers2 = ['السيريال', 'الرمز', 'التاريخ', 'الوقت', 'الحساب', 'النوع', 'قيمة شي ان $',
            'الفاتورة المطابقة', 'قيمة الفاتورة $', 'الحالة', 'ملاحظة']
for ci, h in enumerate(headers2, 1):
    c = ws2.cell(1, ci, h)
    c.font = header_font
    c.fill = header_fill
ws2.freeze_panes = 'A2'
row_i = 2
rows_to_show = [r for r in all_detail_rows if r['card'] in unmatched_card_ids]
for r in rows_to_show:
    ws2.cell(row_i, 1, r['card'])
    ws2.cell(row_i, 2, r.get('code'))
    ws2.cell(row_i, 3, str(r['date']))
    ws2.cell(row_i, 4, str(r['time']))
    ws2.cell(row_i, 5, r['account'])
    ws2.cell(row_i, 6, r['type'])
    ws2.cell(row_i, 7, r['shein_amount'])
    ws2.cell(row_i, 8, r['invoice'])
    ws2.cell(row_i, 9, r['invoice_value'])
    ws2.cell(row_i, 10, r['status'])
    ws2.cell(row_i, 11, r['note'])
    fill = None
    if '❌' in r['status'] or '🟠' in r['status']:
        fill = red
    elif '⚠' in r['status']:
        fill = orange
    elif '✅' in r['status']:
        fill = green
    elif r['status'] == 'تفعيل':
        fill = grey
    if fill:
        for ci in range(1, 12):
            ws2.cell(row_i, ci).fill = fill
    row_i += 1
for ci, w in zip(range(1, 12), [22, 10, 12, 9, 26, 10, 12, 20, 12, 44, 60]):
    ws2.column_dimensions[chr(64 + ci)].width = w

ws3 = wb.create_sheet('الكروت المربوطة')
ws3.sheet_view.rightToLeft = True
headers3 = ['السيريال', 'الرمز', 'الحساب السادس (Shopping Account)', 'ملاحظة']
for ci, h in enumerate(headers3, 1):
    c = ws3.cell(1, ci, h)
    c.font = header_font
    c.fill = header_fill
for ri, lc in enumerate(LINKED_CARDS, 2):
    ws3.cell(ri, 1, lc['card'])
    ws3.cell(ri, 2, code_by_serial.get(lc['card']))
    ws3.cell(ri, 3, lc['shopping_account'])
    ws3.cell(ri, 4, lc['note'])
    for ci in range(1, 5):
        ws3.cell(ri, ci).fill = orange
for ci, w in zip(range(1, 5), [22, 10, 40, 80]):
    ws3.column_dimensions[chr(64 + ci)].width = w

wb.save(out_path)
print()
print('تم الحفظ في:', out_path)
print(f"إجمالي الصفوف التفصيلية: {len(all_detail_rows)}  |  إجمالي البطاقات (بيانات كاملة): {len(all_summaries)}  |  كروت مربوطة: {len(LINKED_CARDS)}")
