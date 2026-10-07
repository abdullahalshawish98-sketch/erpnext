"""
مراجعة تفصيلية كاملة (كل حركة شي ان مقابل كل فاتورة شراء) لكل بطاقات متجر
سلتك، بنفس منطق المطابقة المُثبَت في reconcile_gift_card_movements.py
(مطابقة مباشرة -> إقران استرجاع ضمن 10-180 دقيقة -> إقران موسَّع حتى 24
ساعة لا يُقبَل إلا بتأكيد فاتورة فعلية)، لكنه يَطبع/يُخرِج **كل** حركة
(مطابقة أو غير مطابقة)، لا فقط غير المفسَّرة، حتى يمكن تدقيق أبسط التفاصيل
سطراً بسطر. قراءة فقط — لا يُعدِّل أي بيانات في ERPNext.

الاستخدام: python scripts/detailed_card_review.py <مسار ملف حركات شي ان xlsx> [مسار ملف الإخراج xlsx]
"""
import sys
import openpyxl
import json
import re
from datetime import datetime
from collections import defaultdict
from openpyxl.styles import Font, PatternFill, Alignment

sys.path.insert(0, '.')
import agent_map

shein_file = sys.argv[1] if len(sys.argv) > 1 else 'data/shein_all_movements.xlsx'
out_path = sys.argv[2] if len(sys.argv) > 2 else 'data/detailed_card_review.xlsx'

wb_in = openpyxl.load_workbook(shein_file, data_only=True)
ws_all = wb_in['All Movements']
rows_all = list(ws_all.iter_rows(values_only=True))[1:]

shein_cards = {}
for r in rows_all:
    cid = str(r[2])
    shein_cards.setdefault(cid, []).append(r)

with open('data/Gift_Card.json', encoding='utf-8') as f:
    cards = json.load(f)
with open('data/Purchase_Invoice.json', encoding='utf-8') as f:
    pis = json.load(f)


def base_serial(name):
    m = re.search(r'(\d{10,})', name)
    return m.group(1) if m else None


card_by_serial = {}
card_record_by_serial = {}
for c in cards:
    bs = base_serial(c['name'])
    if bs:
        card_by_serial[bs] = c['name']
        card_record_by_serial[bs] = c


def dt(date_, time_):
    return datetime.fromisoformat(f"{date_} {time_}")


def take(erp_by_value, value):
    v = round(value, 2)
    names = erp_by_value.get(v)
    if names:
        return names.pop()
    return None


def candidates(purchases, refunds, p_idx, r_idx, lo_min, hi_min, target_min):
    cands = []
    for pi in p_idx:
        for ri in r_idx:
            delta_min = (dt(refunds[ri]['date'], refunds[ri]['time']) - dt(purchases[pi]['date'], purchases[pi]['time'])).total_seconds() / 60
            if lo_min <= delta_min <= hi_min:
                cands.append((abs(delta_min - target_min), pi, ri))
    cands.sort(key=lambda x: x[0])
    return cands


def agent_of(cid):
    name = card_by_serial.get(cid)
    if not name:
        return '?'
    letter = name.split(' ', 1)[0]
    return agent_map.CARD_LETTER_TO_AGENT.get(letter) or agent_map.OFFICE_ACCOUNTS_BY_LETTER.get(letter) or letter


# ---------------------------------------------------------------------------
# معالجة بطاقة واحدة: تُعيد (قائمة صفوف تفصيلية، ملخص البطاقة)
# ---------------------------------------------------------------------------
def process_card(cid):
    name = card_by_serial.get(cid)
    erp_rows = sorted([p for p in pis if p.get('gift_card') == name and p.get('docstatus') == 1],
                       key=lambda p: (p.get('posting_date', ''), p.get('creation', '')))
    erp_by_value = defaultdict(list)
    for p in erp_rows:
        erp_by_value[round(p.get('total') or 0, 2)].append(p.get('name'))
    return_names = {p['name'] for p in erp_rows if p.get('is_return')}
    all_invoice_names = {p.get('name') for p in erp_rows}

    detail_rows = []
    used_invoice_names = set()

    events_all = shein_cards[cid]
    activated = [t for t in events_all if t[7] == 'Activated']
    for t in activated:
        detail_rows.append({
            'card': cid, 'seq': t[3], 'date': t[4], 'time': t[5], 'account': t[6],
            'type': 'تفعيل', 'shein_amount': t[8], 'invoice': '', 'invoice_value': '',
            'status': 'تفعيل', 'note': '',
        })

    accounts = sorted({t[6] for t in events_all if t[7] != 'Activated'})
    acc_all_purchase_results = defaultdict(list)
    acc_leftover_refund_events = defaultdict(list)

    for acc in accounts:
        events = sorted((t for t in events_all if t[7] != 'Activated' and t[6] == acc), key=lambda t: (t[4], t[5]))
        sessions = []
        cur, prev_dt = [], None
        for t in events:
            cur_dt = dt(t[4], t[5])
            if prev_dt and (cur_dt - prev_dt).total_seconds() > 3 * 3600:
                sessions.append(cur)
                cur = []
            cur.append(t)
            prev_dt = cur_dt
        if cur:
            sessions.append(cur)

        acc_leftover_purchases = []
        acc_leftover_refunds = []

        for sess in sessions:
            purchases = [{'amount': -t[8], 'date': t[4], 'time': t[5], 'account': acc, 'seq': t[3]} for t in sess if t[7] == 'Purchase']
            refunds = [{'amount': t[8], 'date': t[4], 'time': t[5], 'account': acc, 'seq': t[3]} for t in sess if t[7] == 'Refund']

            if purchases and refunds and abs(sum(p['amount'] for p in purchases) - sum(r['amount'] for r in refunds)) < 0.05:
                for p in purchases:
                    detail_rows.append({
                        'card': cid, 'seq': p['seq'], 'date': p['date'], 'time': p['time'], 'account': acc,
                        'type': 'شراء', 'shein_amount': -p['amount'], 'invoice': '', 'invoice_value': '',
                        'status': '✅ جلسة مُلغاة بالكامل', 'note': 'مجموع الشراء = مجموع الاسترجاع (قطعة توصيل أُضيفت وأُلغيت)',
                    })
                for r in refunds:
                    detail_rows.append({
                        'card': cid, 'seq': r['seq'], 'date': r['date'], 'time': r['time'], 'account': acc,
                        'type': 'استرجاع', 'shein_amount': r['amount'], 'invoice': '', 'invoice_value': '',
                        'status': '✅ جلسة مُلغاة بالكامل', 'note': 'ضمن إلغاء الجلسة الكاملة',
                    })
                acc_all_purchase_results[acc].extend({**p, 'net': 0.0} for p in purchases)
                continue

            all_p, all_r = list(range(len(purchases))), list(range(len(refunds)))
            used_p, used_r = set(), set()
            note_of = {}

            for pi in all_p:
                matched_name = take(erp_by_value, purchases[pi]['amount'])
                if matched_name:
                    used_p.add(pi)
                    note_of[pi] = ('مطابقة مباشرة بقيمة شي ان الكاملة', purchases[pi]['amount'], matched_name)

            remaining_p = [i for i in all_p if i not in used_p]
            for _, pi, ri in candidates(purchases, refunds, remaining_p, all_r, 10, 180, 60):
                if pi in used_p or ri in used_r:
                    continue
                net = round(purchases[pi]['amount'] - refunds[ri]['amount'], 2)
                matched_name = take(erp_by_value, net)
                used_p.add(pi)
                used_r.add(ri)
                gap = (dt(refunds[ri]['date'], refunds[ri]['time']) - dt(purchases[pi]['date'], purchases[pi]['time'])).total_seconds() / 60
                note_of[pi] = (f"مقرون باسترجاع +{refunds[ri]['amount']}$ بعد {gap:.0f} دقيقة", net, matched_name)
                detail_rows.append({
                    'card': cid, 'seq': refunds[ri]['seq'], 'date': refunds[ri]['date'], 'time': refunds[ri]['time'], 'account': acc,
                    'type': 'استرجاع', 'shein_amount': refunds[ri]['amount'],
                    'invoice': matched_name or '', 'invoice_value': '',
                    'status': '✅ مطروح من فاتورة الشراء المقابلة' if matched_name else '⚠ مقرون زمنياً لكن القيمة الصافية غير موجودة كفاتورة',
                    'note': f"نُسخ ضمن صافي الشراء {purchases[pi]['date']} {purchases[pi]['time']}",
                })

            for pi, p in enumerate(purchases):
                if pi in note_of:
                    note, net, matched_name = note_of[pi]
                else:
                    note, net, matched_name = 'بلا استرجاع (بعد)', round(p['amount'], 2), None
                p_result = {**p, 'net': net, 'note': note, 'match': matched_name}
                acc_all_purchase_results[acc].append(p_result)
                if pi not in note_of:
                    acc_leftover_purchases.append(p_result)

            for ri in all_r:
                if ri not in used_r:
                    acc_leftover_refunds.append(refunds[ri])

        # إقران موسَّع حتى 24 ساعة على مستوى الحساب كله
        all_p2, all_r2 = list(range(len(acc_leftover_purchases))), list(range(len(acc_leftover_refunds)))
        used_p2, used_r2 = set(), set()
        for _, pi, ri in candidates(acc_leftover_purchases, acc_leftover_refunds, all_p2, all_r2, 180, 1440, 180):
            if pi in used_p2 or ri in used_r2:
                continue
            net = round(acc_leftover_purchases[pi]['net'] - acc_leftover_refunds[ri]['amount'], 2)
            matched_name = take(erp_by_value, net)
            if matched_name:
                used_p2.add(pi)
                used_r2.add(ri)
                r = acc_leftover_refunds[ri]
                gap = (dt(r['date'], r['time']) - dt(acc_leftover_purchases[pi]['date'], acc_leftover_purchases[pi]['time'])).total_seconds() / 60
                acc_leftover_purchases[pi]['net'] = net
                acc_leftover_purchases[pi]['note'] = f"⚠ شاذة — مقرون باسترجاع متأخر {r['date']} {r['time']} بعد {gap:.0f} دقيقة (نسيان إلغاء محتمل)"
                acc_leftover_purchases[pi]['match'] = matched_name
                detail_rows.append({
                    'card': cid, 'seq': r['seq'], 'date': r['date'], 'time': r['time'], 'account': acc,
                    'type': 'استرجاع', 'shein_amount': r['amount'], 'invoice': matched_name or '', 'invoice_value': '',
                    'status': '✅ مطروح (إقران موسَّع حتى 24 ساعة)',
                    'note': f"نُسخ ضمن صافي الشراء {acc_leftover_purchases[pi]['date']} {acc_leftover_purchases[pi]['time']}",
                })

        for r_idx2, r in enumerate(acc_leftover_refunds):
            if r_idx2 in used_r2:
                continue
            acc_leftover_refund_events[acc].append(r)

        # صفوف الشراء التفصيلية لهذا الحساب
        for p in acc_all_purchase_results[acc]:
            if 'match' not in p:
                continue
            match = p.get('match')
            net = p['net']
            if match:
                used_invoice_names.add(match)
                inv_val = next((pp.get('total') for pp in erp_rows if pp.get('name') == match), '')
                status = f'✅ مطابق {match}'
            else:
                status = '❌ غير مفسَّر — لا توجد فاتورة بهذه القيمة الصافية'
                inv_val = ''
            detail_rows.append({
                'card': cid, 'seq': p['seq'], 'date': p['date'], 'time': p['time'], 'account': p['account'],
                'type': 'شراء', 'shein_amount': -p['amount'] if 'amount' in p else '',
                'invoice': match or '', 'invoice_value': inv_val,
                'status': status, 'note': f"{p['note']} | صافي متوقَّع={net}$",
            })

        # استرجاعات باقية بلا أي شراء (تشخيص نسيان طرح / نفاذ مخزون)
        for lr in acc_leftover_refund_events.get(acc, []):
            r_dt = dt(lr['date'], lr['time'])
            best = None
            for p in acc_all_purchase_results[acc]:
                p_dt = dt(p['date'], p['time'])
                if p_dt >= r_dt:
                    continue
                gap_min = (r_dt - p_dt).total_seconds() / 60
                if best is None or gap_min < best[0]:
                    best = (gap_min, p)
            if best:
                gap_min, p = best
                guess = 'نفس الفاتورة القريبة (نسيان طرح محتمل)' if gap_min <= 180 else '📦 نفاذ مخزون محتمل — فاتورة مرتجع منفصلة'
                note = f"أقرب شراء سابق {p['date']} {p['time']} صافيه={p.get('net')}$ (فاتورة {p.get('match')}) | فارق {gap_min:.0f} دقيقة -> {guess}"
            else:
                note = 'لا يوجد شراء سابق على نفس الحساب'
            detail_rows.append({
                'card': cid, 'seq': lr['seq'], 'date': lr['date'], 'time': lr['time'], 'account': lr['account'],
                'type': 'استرجاع', 'shein_amount': lr['amount'], 'invoice': '', 'invoice_value': '',
                'status': '⚠ استرجاع غير مُستخدَم', 'note': note,
            })

    # فواتير ERPNext لهذه البطاقة لم تُستخدَم في أي مطابقة إطلاقاً
    unused_invoices = all_invoice_names - used_invoice_names
    for inv_name in unused_invoices:
        inv = next(p for p in erp_rows if p.get('name') == inv_name)
        detail_rows.append({
            'card': cid, 'seq': '', 'date': inv.get('posting_date'), 'time': str(inv.get('posting_time'))[:8],
            'account': '', 'type': 'فاتورة مرتجع' if inv.get('is_return') else 'فاتورة شراء',
            'shein_amount': '', 'invoice': inv_name, 'invoice_value': inv.get('total'),
            'status': '🟠 فاتورة يتيمة — لا تقابلها أي حركة شي ان مطابقة', 'note': '',
        })

    detail_rows.sort(key=lambda r: (str(r['date']), str(r['time'])))

    # ملخص البطاقة
    real_balance = card_record_by_serial.get(cid, {}).get('balance')
    amount = card_record_by_serial.get(cid, {}).get('amount')
    purchases_sum = sum(-t[8] for t in events_all if t[7] == 'Purchase')
    refunds_sum = sum(t[8] for t in events_all if t[7] == 'Refund')
    shein_real_balance = round((activated[0][8] if activated else 0) - purchases_sum + refunds_sum, 2)
    unmatched_purchases = sum(1 for r in detail_rows if r['type'] == 'شراء' and '❌' in r['status'])
    unused_refunds = sum(1 for r in detail_rows if r['status'] == '⚠ استرجاع غير مُستخدَم')
    orphan_invoices = sum(1 for r in detail_rows if 'يتيمة' in r['status'])

    summary = {
        'card': cid, 'name': name, 'agent': agent_of(cid),
        'activated': activated[0][8] if activated else '',
        'erpnext_balance': real_balance,
        'shein_real_balance': shein_real_balance,
        'diff': round((real_balance or 0) - shein_real_balance, 2) if real_balance is not None else '',
        'unmatched_purchases': unmatched_purchases,
        'unused_refunds': unused_refunds,
        'orphan_invoices': orphan_invoices,
    }
    return detail_rows, summary


all_detail_rows = []
all_summaries = []
for cid in shein_cards:
    if cid not in card_by_serial:
        print(f"!! تحذير: لا توجد بطاقة في ERPNext للسيريال {cid}")
        continue
    detail_rows, summary = process_card(cid)
    all_detail_rows.extend(detail_rows)
    all_summaries.append(summary)
    flag = '  <<< يحتاج مراجعة' if (summary['unmatched_purchases'] or summary['unused_refunds'] or summary['orphan_invoices'] or abs(summary['diff'] or 0) > 0.05) else ''
    print(f"{cid}  ({summary['agent']})  erp={summary['erpnext_balance']}  shein={summary['shein_real_balance']}  diff={summary['diff']}  "
          f"غير مطابق={summary['unmatched_purchases']}  استرجاع غير مستخدم={summary['unused_refunds']}  فواتير يتيمة={summary['orphan_invoices']}{flag}")

all_summaries.sort(key=lambda s: abs(s['diff'] if isinstance(s['diff'], (int, float)) else 0), reverse=True)

# ---------------------------------------------------------------------------
# بناء ملف Excel
# ---------------------------------------------------------------------------
wb = openpyxl.Workbook()
bold = Font(bold=True)
red = PatternFill('solid', fgColor='FFC7CE')
orange = PatternFill('solid', fgColor='FFE0B2')
green = PatternFill('solid', fgColor='C6EFCE')
grey = PatternFill('solid', fgColor='E0E0E0')
header_fill = PatternFill('solid', fgColor='4472C4')
header_font = Font(bold=True, color='FFFFFF')

ws1 = wb.active
ws1.title = 'ملخص كل الكروت'
ws1.sheet_view.rightToLeft = True
headers1 = ['السيريال', 'اسم البطاقة', 'الوكيل', 'التفعيل $', 'رصيد المنظومة $', 'الرصيد الحقيقي (شي ان) $',
            'الفرق $', 'مشتريات غير مفسَّرة', 'استرجاعات غير مستخدَمة', 'فواتير يتيمة']
for ci, h in enumerate(headers1, 1):
    c = ws1.cell(1, ci, h)
    c.font = header_font
    c.fill = header_fill
for ri, s in enumerate(all_summaries, 2):
    ws1.cell(ri, 1, s['card'])
    ws1.cell(ri, 2, s['name'])
    ws1.cell(ri, 3, s['agent'])
    ws1.cell(ri, 4, s['activated'])
    ws1.cell(ri, 5, s['erpnext_balance'])
    ws1.cell(ri, 6, s['shein_real_balance'])
    ws1.cell(ri, 7, s['diff'])
    ws1.cell(ri, 8, s['unmatched_purchases'])
    ws1.cell(ri, 9, s['unused_refunds'])
    ws1.cell(ri, 10, s['orphan_invoices'])
    is_clean = s['unmatched_purchases'] == 0 and s['unused_refunds'] == 0 and s['orphan_invoices'] == 0 and abs(s['diff'] or 0) <= 0.05
    fill = green if is_clean else (red if abs(s['diff'] or 0) > 0.05 else orange)
    for ci in range(1, 11):
        ws1.cell(ri, ci).fill = fill
for ci, w in zip(range(1, 11), [22, 34, 22, 12, 16, 20, 10, 14, 16, 12]):
    ws1.column_dimensions[chr(64 + ci)].width = w

ws2 = wb.create_sheet('تفاصيل كل الحركات')
ws2.sheet_view.rightToLeft = True
headers2 = ['السيريال', 'الوكيل', '#', 'التاريخ', 'الوقت', 'الحساب', 'النوع', 'قيمة شي ان $',
            'الفاتورة المطابقة', 'قيمة الفاتورة $', 'الحالة', 'ملاحظة']
for ci, h in enumerate(headers2, 1):
    c = ws2.cell(1, ci, h)
    c.font = header_font
    c.fill = header_fill
ws2.freeze_panes = 'A2'
row_i = 2
last_card = None
for r in all_detail_rows:
    if r['card'] != last_card:
        last_card = r['card']
    ws2.cell(row_i, 1, r['card'])
    ws2.cell(row_i, 2, agent_of(r['card']))
    ws2.cell(row_i, 3, r['seq'])
    ws2.cell(row_i, 4, str(r['date']))
    ws2.cell(row_i, 5, str(r['time']))
    ws2.cell(row_i, 6, r['account'])
    ws2.cell(row_i, 7, r['type'])
    ws2.cell(row_i, 8, r['shein_amount'])
    ws2.cell(row_i, 9, r['invoice'])
    ws2.cell(row_i, 10, r['invoice_value'])
    ws2.cell(row_i, 11, r['status'])
    ws2.cell(row_i, 12, r['note'])
    fill = None
    if '❌' in r['status'] or 'يتيمة' in r['status']:
        fill = red
    elif '⚠' in r['status']:
        fill = orange
    elif '✅' in r['status']:
        fill = green
    elif r['status'] == 'تفعيل':
        fill = grey
    if fill:
        for ci in range(1, 13):
            ws2.cell(row_i, ci).fill = fill
    row_i += 1
for ci, w in zip(range(1, 13), [22, 18, 5, 12, 9, 26, 10, 12, 20, 12, 44, 60]):
    ws2.column_dimensions[chr(64 + ci)].width = w

wb.save(out_path)
print()
print('تم الحفظ في:', out_path)
print(f"إجمالي الصفوف التفصيلية: {len(all_detail_rows)}  |  إجمالي البطاقات: {len(all_summaries)}")
